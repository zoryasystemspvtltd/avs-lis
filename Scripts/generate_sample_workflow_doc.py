"""Generate Sample Collection & Sample Receiving architecture Word document."""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from datetime import datetime
import os

OUTPUT_PATH = r"I:\Projects\LIS\ZoryaLIS_Sample_Collection_and_Receiving_Guide.docx"


def add_heading(doc, text, level=1):
    return doc.add_heading(text, level=level)


def add_para(doc, text, bold=False, italic=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(11)
    return p


def add_bullet(doc, text, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(text, style=style)
    for run in p.runs:
        run.font.size = Pt(11)
    return p


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        for p in hdr_cells[i].paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(10)
    for ri, row in enumerate(rows):
        cells = table.rows[ri + 1].cells
        for ci, val in enumerate(row):
            cells[ci].text = str(val)
            for p in cells[ci].paragraphs:
                for run in p.runs:
                    run.font.size = Pt(10)
    doc.add_paragraph()
    return table


def build_document():
    doc = Document()

    # Title page
    title = doc.add_heading("ZoryaLIS", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph("Sample Collection & Sample Receiving")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in sub.runs:
        run.bold = True
        run.font.size = Pt(16)
    sub2 = doc.add_paragraph("Data Flow & Business Logic Guide")
    sub2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in sub2.runs:
        run.font.size = Pt(14)
        run.italic = True
    meta = doc.add_paragraph(f"Generated: {datetime.now().strftime('%d %B %Y')}")
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()

    # 1. Overview
    add_heading(doc, "1. Overview", 1)
    add_para(
        doc,
        "Sample Collection and Sample Receiving are the pre-analytical workflow stages in ZoryaLIS. "
        "They move a laboratory order from \"paid and waiting\" through \"physically collected\" to "
        "\"received in the laboratory.\" Both stages operate on the core entity TestRequestDetails "
        "in the AVSLIS database."
    )

    # 2. End-to-end flow
    add_heading(doc, "2. End-to-End Data Flow", 1)
    add_para(doc, "The workflow follows this sequence:", bold=True)
    add_bullet(doc, "Patient Master / Sale Invoice — order and test request creation")
    add_bullet(doc, "Sample Collection — phlebotomist collects specimen (paid invoices only)")
    add_bullet(doc, "Sample Receiving — lab reception accepts the sample")
    add_bullet(doc, "Downstream — analyzer results, Edit Results, technician/doctor approval, reports")

    add_para(doc, "High-level flow:", bold=True)
    flow_steps = [
        "Sale Invoice saved with lab test lines",
        "TestRequestDetails record created (ReportStatus = New)",
        "Invoice marked Paid → appears in Collection queue",
        "Collect → CollectedBy and SampleCollectionDate set",
        "Receive → ReceivedBy and SampleReceivedDate set",
        "Results pipeline → TestResults, approval, final report",
    ]
    for i, step in enumerate(flow_steps, 1):
        add_bullet(doc, f"{i}. {step}")

    # 3. Order creation
    add_heading(doc, "3. Where Orders Come From (Upstream)", 1)
    add_para(
        doc,
        "The primary entry path today is Sale Invoice (not the legacy New Sample screen). "
        "When a Sale Invoice is saved with laboratory test lines, SaleInvoiceManager.LinkTestRequestsToLines() "
        "creates or reuses a TestRequestDetail record per test."
    )
    add_para(doc, "On creation:", bold=True)
    add_bullet(doc, "Barcode (SampleNo) is pre-generated as {RequestNo}-{TestCode} (e.g. INV-20260519-0002-CBC)")
    add_bullet(doc, "ReportStatus is set to New (0)")
    add_bullet(doc, "CollectedBy and ReceivedBy are empty — these drive workflow gates")
    add_bullet(doc, "Radiology tests route to RadiologyRequestDetail, not TestRequestDetails")

    add_table(
        doc,
        ["Field", "Initial Value", "Purpose"],
        [
            ["SampleNo", "{RequestNo}-{TestCode}", "Barcode label / tracking"],
            ["HISRequestNo", "Invoice request number", "Order reference"],
            ["ReportStatus", "New (0)", "Overall report lifecycle"],
            ["CollectedBy", "Empty", "Collection gate"],
            ["ReceivedBy", "Empty", "Receiving gate"],
            ["PatientId", "From invoice", "Patient linkage"],
        ],
    )

    # 4. State model
    add_heading(doc, "4. State Model (Business Logic Core)", 1)
    add_para(
        doc,
        "Workflow progress is tracked on TestRequestDetails using ReportStatus, CollectedBy, and ReceivedBy."
    )
    add_table(
        doc,
        ["ReportStatus Value", "Name", "Typical Meaning"],
        [
            ["0", "New", "Awaiting collection / receiving / results"],
            ["1", "SentToEquipment", "Dispatched to analyzer"],
            ["2", "ReportGenerated", "Results entered or received from equipment"],
            ["3", "TechnicianApproved", "Technician signed off"],
            ["4", "TechnicianRejected", "Technician rejected"],
            ["5", "DoctorApproved", "Doctor authorized report"],
            ["6", "DoctorRejected", "Doctor rejected"],
            ["7", "FinallyRejected", "Sample rejected at collection or receiving"],
        ],
    )
    add_table(
        doc,
        ["Field", "Empty", "Set"],
        [
            ["CollectedBy", "Not collected", "Sample collected by phlebotomist"],
            ["ReceivedBy", "Not received in lab", "Sample received at lab reception"],
        ],
    )

    # 5. Sample Collection
    add_heading(doc, "5. Sample Collection", 1)
    add_table(
        doc,
        ["Layer", "Component"],
        [
            ["Portal route", "/sample-collection"],
            ["Angular component", "SampleCollectionComponent"],
            ["API prefix", "api/SampleCollection"],
            ["Business layer", "SampleCollectionManager"],
            ["Database table", "TestRequestDetails"],
        ],
    )

    add_heading(doc, "5.1 Pending Collection Queue", 2)
    add_para(doc, "A row appears in the queue only when ALL of the following are true:", bold=True)
    add_bullet(doc, "ReportStatus = New")
    add_bullet(doc, "CollectedBy is empty (not yet collected)")
    add_bullet(doc, "The request is linked to a Paid invoice (InvoiceStatus = Paid)")
    add_para(
        doc,
        "Important: Saving an invoice creates test requests, but collection waits until payment. "
        "The paid-invoice link is resolved via SaleInvoiceDetail.RequestDetailId or matching HISRequestNo."
    )

    add_heading(doc, "5.2 Collect Action", 2)
    add_para(doc, "API: POST api/SampleCollection/Collect", bold=True)
    add_para(doc, "Business rules:", bold=True)
    add_bullet(doc, "Duplicate collection is not allowed")
    add_bullet(doc, "Collection date/time cannot be in the future")
    add_bullet(doc, "Barcode must be unique across orders")
    add_bullet(doc, "Updates: SampleNo, SampleCollectionDate, CollectedBy, CollectedRemarks")
    add_para(
        doc,
        "After collection, the row leaves the collection queue and becomes eligible for Sample Receiving."
    )

    add_heading(doc, "5.3 Reject at Collection", 2)
    add_para(doc, "API: POST api/SampleCollection/Reject")
    add_bullet(doc, "Sets ReportStatus = FinallyRejected")
    add_bullet(doc, "Stores rejection reason in CollectedRemarks")
    add_bullet(doc, "Sample is removed from normal workflow")

    add_heading(doc, "5.4 Recollect at Collection", 2)
    add_para(doc, "API: POST api/SampleCollection/Recollect/{id}")
    add_bullet(doc, "Calls TestRequestDetailsManager.TechnicianReview with ReportStatusType.New")
    add_bullet(doc, "Marks old request as FinallyRejected")
    add_bullet(doc, "Creates a new TestRequestDetail with ReportStatus = New for a fresh collection cycle")

    # 6. Sample Receiving
    add_heading(doc, "6. Sample Receiving", 1)
    add_table(
        doc,
        ["Layer", "Component"],
        [
            ["Portal route", "/sample-receiving"],
            ["Angular component", "SampleReceivingComponent"],
            ["API prefix", "api/SampleReceiving"],
            ["Business layer", "SampleReceivingManager"],
            ["Database table", "TestRequestDetails"],
        ],
    )

    add_heading(doc, "6.1 Receiving Queue", 2)
    add_para(doc, "A row appears when:", bold=True)
    add_bullet(doc, "CollectedBy is set (sample was collected)")
    add_bullet(doc, "ReceivedBy is empty (not yet received)")
    add_bullet(doc, "ReportStatus = New")
    add_para(
        doc,
        "There is no paid-invoice filter at receiving — if the sample was collected, it can be received."
    )

    add_heading(doc, "6.2 Receive Action", 2)
    add_para(doc, "API: POST api/SampleReceiving/Receive", bold=True)
    add_para(doc, "Business rules:", bold=True)
    add_bullet(doc, "Sample must be collected before receiving")
    add_bullet(doc, "Duplicate receiving is not allowed")
    add_bullet(doc, "Receiving time cannot be before collection time")
    add_bullet(doc, "Barcode in request must match the order (if provided)")
    add_bullet(doc, "Updates: SampleReceivedDate, ReceivedBy, ReceivedRemarks")

    add_heading(doc, "6.3 Reject at Receiving", 2)
    add_para(doc, "API: POST api/SampleReceiving/Reject", bold=True)
    add_bullet(doc, "Rejection reason code is mandatory (from SampleRejectionReasonMaster, category \"Receiving\")")
    add_bullet(doc, "Sets ReportStatus = FinallyRejected")
    add_bullet(doc, "Stores reason in ReceivedRemarks")

    add_heading(doc, "6.4 Recollect at Receiving", 2)
    add_para(doc, "Same recollection pattern as collection — creates a new test request cycle.")

    # 7. UI/API map
    add_heading(doc, "7. UI → API → Business Layer Map", 1)
    add_table(
        doc,
        ["User Action", "Angular Service", "API Endpoint", "Manager Method"],
        [
            ["Load collection queue", "getCollectionQueue()", "GET SampleCollection/PendingQueue", "GetPendingQueue()"],
            ["Collect sample", "collectSample()", "POST SampleCollection/Collect", "CollectSample()"],
            ["Reject (collection)", "rejectCollection()", "POST SampleCollection/Reject", "RejectCollection()"],
            ["Recollect", "recollectCollection()", "POST SampleCollection/Recollect/{id}", "TriggerRecollection()"],
            ["Load receiving queue", "getReceivingQueue()", "GET SampleReceiving/Queue", "GetReceivingQueue()"],
            ["Receive sample", "receiveSample()", "POST SampleReceiving/Receive", "ReceiveSample()"],
            ["Reject (receiving)", "rejectReceiving()", "POST SampleReceiving/Reject", "RejectSample()"],
            ["Rejection reasons", "getRejectionReasons()", "GET SampleReceiving/RejectionReasons", "(repository read)"],
        ],
    )
    add_para(
        doc,
        "Search filters (barcode, order number, patient name, etc.) are passed via the ApiOption HTTP header "
        "as JSON — consistent with other ZoryaLIS list screens. Both screens support bulk collect/receive "
        "(one API call per selected row via forkJoin on the frontend)."
    )

    # 8. After receiving
    add_heading(doc, "8. What Happens After Receiving", 1)
    add_bullet(doc, "Equipment / Analyzer — results ingested via ResultManager (creates TestResult + TestResultDetails)")
    add_bullet(doc, "Edit Results — manual result entry when needed")
    add_bullet(doc, "Technician / Doctor review — approval via TestRequestDetailsManager")
    add_bullet(doc, "Reports — TAT metrics use SampleCollectionDate and SampleReceivedDate")
    add_para(
        doc,
        "Receiving does not itself create result rows — it only marks the physical handoff into the laboratory."
    )

    # 9. Lifecycle
    add_heading(doc, "9. Sample Lifecycle (Single Specimen)", 1)
    lifecycle = [
        ("Order Created", "Sale Invoice saved → TestRequestDetails created"),
        ("Pending Collection", "Invoice Paid → appears in Collection queue"),
        ("Collected", "Sample Collection → CollectedBy set"),
        ("Awaiting Receiving", "In transit / at lab desk"),
        ("Received", "Sample Receiving → ReceivedBy set"),
        ("Results Pipeline", "Analyzer / Edit Results"),
        ("Approved", "Technician + Doctor sign-off"),
        ("Rejected", "Collection or Receiving reject → FinallyRejected"),
    ]
    add_table(doc, ["Stage", "Description"], lifecycle)

    # 10. Permissions
    add_heading(doc, "10. Permissions & Modules", 1)
    add_table(
        doc,
        ["Module", "Permissions", "Menu Location"],
        [
            ["SampleCollection", "View, Edit, Reject", "LIS workflow left nav"],
            ["SampleReceiving", "View, Edit, Reject", "LIS workflow left nav"],
        ],
    )

    # 11. Practical summary
    add_heading(doc, "11. Practical Summary", 1)
    add_table(
        doc,
        ["Stage", "Who", "Queue Condition", "DB Fields Updated"],
        [
            ["Order", "Billing / front desk", "—", "TestRequestDetails created from Sale Invoice"],
            ["Collection", "Phlebotomist", "Paid + not collected", "CollectedBy, SampleCollectionDate, SampleNo, CollectedRemarks"],
            ["Receiving", "Lab reception", "Collected + not received", "ReceivedBy, SampleReceivedDate, ReceivedRemarks"],
            ["Results", "Lab tech / analyzer", "Received", "TestResults, ReportStatus progression"],
        ],
    )

    # 12. Key source files
    add_heading(doc, "12. Key Source Files (Reference)", 1)
    add_table(
        doc,
        ["Purpose", "File Path"],
        [
            ["Collection business logic", "LIS.Businesslogic\\SampleCollectionManager.cs"],
            ["Receiving business logic", "LIS.Businesslogic\\SampleReceivingManager.cs"],
            ["Test request / recollection", "LIS.Businesslogic\\TestRequestDetailsManager.cs"],
            ["Invoice → test request creation", "LIS.Businesslogic\\SaleInvoiceManager.cs"],
            ["Collection API", "web\\Lis.Api\\Controllers\\Api\\SampleCollectionController.cs"],
            ["Receiving API", "web\\Lis.Api\\Controllers\\Api\\SampleReceivingController.cs"],
            ["Collection UI", "web\\Lis.Web\\src\\app\\LIS\\sample-workflow\\sample-collection\\"],
            ["Receiving UI", "web\\Lis.Web\\src\\app\\LIS\\sample-workflow\\sample-receiving\\"],
            ["Angular workflow service", "web\\Lis.Web\\src\\app\\_services\\sample-workflow.service.ts"],
            ["Entity model", "LIS.DtoModel\\Models\\DTO\\Inflow\\TestRequestDetails.cs"],
        ],
    )

    # Footer
    doc.add_paragraph()
    footer = doc.add_paragraph(
        "Document produced from ZoryaLIS (AVILIS) codebase analysis. "
        "For internal technical reference."
    )
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in footer.runs:
        run.italic = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    return doc


def main():
    os.makedirs(os.path.dirname(OUTPUT_PATH), exist_ok=True)
    doc = build_document()
    doc.save(OUTPUT_PATH)
    print(f"Created: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
