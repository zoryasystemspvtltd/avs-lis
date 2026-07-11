# -*- coding: utf-8 -*-
"""Generate I:\\Projects\\LIS\\DBRemoval.docx — analysis only."""
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from datetime import datetime
import os

doc = Document()
for s in doc.sections:
    s.top_margin = Inches(0.75)
    s.bottom_margin = Inches(0.75)
    s.left_margin = Inches(0.85)
    s.right_margin = Inches(0.85)


def h(text, level=1):
    doc.add_heading(text, level=level)


def p(text, bold=False):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.bold = bold
    run.font.size = Pt(10)
    return para


def bullets(items):
    for i in items:
        para = doc.add_paragraph(i, style="List Bullet")
        for run in para.runs:
            run.font.size = Pt(10)


def add_table(headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    hdr = t.rows[0].cells
    for i, htxt in enumerate(headers):
        hdr[i].text = htxt
        for para in hdr[i].paragraphs:
            for run in para.runs:
                run.bold = True
                run.font.size = Pt(8)
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            t.rows[r_i + 1].cells[c_i].text = str(val)
            for para in t.rows[r_i + 1].cells[c_i].paragraphs:
                for run in para.runs:
                    run.font.size = Pt(8)
    doc.add_paragraph()


title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("ZoryaLMS / ZoryaLIS")
r.bold = True
r.font.size = Pt(22)

st = doc.add_paragraph()
st.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = st.add_run("Enterprise Database Object Usage Analysis & Cleanup Report")
r.bold = True
r.font.size = Pt(16)

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.add_run(
    f"Analysis Date: {datetime.now().strftime('%Y-%m-%d %H:%M')}\n"
    f"Database: ZoryaLMS (SQL Server Express)\n"
    f"Solution: I:\\Projects\\LIS\\avs-lis\n"
    f"Scope: ANALYSIS ONLY — No objects deleted, modified, renamed, or dropped.\n"
    f"Classification standard: SAFE TO REMOVE only with verifiable evidence across "
    f"EF, APIs, UI, reports, scripts, and tests."
).font.size = Pt(10)

doc.add_page_break()

h("Section 1 — Executive Summary")
p(
    "ZoryaLMS is an Entity Framework 6 (Code First) + ASP.NET Identity application. "
    "The live database contains 45 user tables and zero application views, stored procedures, "
    "functions, triggers, sequences, synonyms, or user-defined types. Data access is almost entirely "
    "LINQ/DbSet via ApplicationDBContext (domain) and IdentityDbContext (security). No Dapper usage "
    "and no application stored-procedure calls were found in C#."
)
p("Key findings:", bold=True)
bullets(
    [
        "45 tables in dbo — all map to EF Identity or ApplicationDBContext entities (or EF MigrationHistory).",
        "0 views / 0 procedures / 0 functions / 0 triggers in live ZoryaLMS — SP/View/Function unused matrices are empty for live DB.",
        "Legacy SQL_Script\\ artifacts (views vw*, HisParamMaster, ExecutionLog) target LISStaging/HIS bridge and are NOT present in ZoryaLMS — orphaned relative to current app.",
        "DTO entities TestResultsHist / TestResultDetailsHist exist in code with [Table] attributes but have NO DbSet and NO matching tables in ZoryaLMS.",
        "RoleMenuPermission is newly introduced (EF + API + Role Edit UI) but currently has 0 rows — Partially Used / newly adopted feature.",
        "GenericRepository.ExecuteSQL / Database.SqlQuery infrastructure has zero business callers — dead API surface (code), not a DB object.",
        "DMV shows several non-clustered indexes with 0 reads since last SQL restart — REQUIRES BUSINESS REVIEW (stats reset on restart; do not drop based on DMV alone).",
        "NO object is classified SAFE TO REMOVE in this report. Live tables are all referenced by EF and/or Identity runtime.",
    ]
)
p(
    "Risk posture: Conservative. Objects used only by certification seeds or with zero rows remain "
    "DO NOT TOUCH or REQUIRES BUSINESS REVIEW until product owners confirm retirement."
)

h("Section 2 — Database Statistics")
add_table(
    ["Object Category", "Count in ZoryaLMS", "Notes"],
    [
        ["Tables (user)", "45", "All dbo"],
        ["Views", "0", "None in live DB"],
        ["Stored Procedures", "0", "None in live DB"],
        ["Functions (FN/IF/TF)", "0", "None in live DB"],
        ["Triggers", "0", "None in live DB"],
        ["Sequences", "0", "None"],
        ["Synonyms", "0", "None"],
        ["User-Defined Types", "0", "None"],
        ["Primary Keys", "45", "One per table"],
        ["Unique Constraints (UQ)", "0", "Uniqueness via unique indexes instead"],
        ["Foreign Keys", "42", "sys.foreign_keys"],
        ["Default Constraints", "22", ""],
        ["Check Constraints", "0", ""],
        ["Non-PK Indexes", "45", "Includes unique/business indexes"],
        ["Schemas (relevant)", "dbo (+ system)", "Business objects in dbo"],
        ["EF Domain DbSets", "31+", "ApplicationDBContext"],
        ["EF Identity custom DbSets", "6", "Plus inherited AspNet* sets"],
        ["Lis.Api Migrations", "4", "Includes bootstrap fresh"],
        ["LIS.DataModel Migrations", "10", "Domain schema evolution"],
    ],
)

h("2.1 Complete Table Inventory", 2)
tables = [
    "AspNetRoles",
    "AspNetUserClaims",
    "AspNetUserLogins",
    "AspNetUserRoles",
    "AspNetUsers",
    "ClientApplication",
    "ContainerMaster",
    "ControlResultDetails",
    "ControlResults",
    "CorporateMaster",
    "Department",
    "EquipmentHeartBeat",
    "EquipmentMaster",
    "HISParameterMaster",
    "HISParameterRangMaster",
    "HISSpecimenMaster",
    "HISTestMaster",
    "MethodMaster",
    "MigrationHistory",
    "PatientDetails",
    "PatientVisit",
    "RadiologyRequestDetail",
    "RadiologyResultDetail",
    "ReferralDoctorMaster",
    "RefreshTokens",
    "RoleMenuPermission",
    "RoleModuleMappings",
    "SaleInvoice",
    "SaleInvoiceDetail",
    "SampleRejectionReasonMaster",
    "SampleTypeMaster",
    "TestCategoryMaster",
    "TestGroupMaster",
    "TestMappingMaster",
    "TestParameterMappingMaster",
    "TestParameters",
    "TestProfileDetail",
    "TestProfileMaster",
    "TestRateMaster",
    "TestRequestDetails",
    "TestResultDetails",
    "TestResults",
    "UnitMaster",
    "UserApplicationMappings",
    "UserModules",
]
p("Alphabetical inventory (45): " + ", ".join(tables))

h("Section 3 — Table Usage Matrix")
p("Classification legend: Used | Partially Used | Legacy | Not Referenced (in live DB) | Unknown")

row_map = {
    "AspNetRoles": 3,
    "AspNetUserClaims": 0,
    "AspNetUserLogins": 0,
    "AspNetUserRoles": 10,
    "AspNetUsers": 8,
    "ClientApplication": 1,
    "ContainerMaster": 49,
    "ControlResultDetails": 3,
    "ControlResults": 2,
    "CorporateMaster": 286,
    "Department": 16,
    "EquipmentHeartBeat": 1,
    "EquipmentMaster": 144,
    "HISParameterMaster": 36,
    "HISParameterRangMaster": 34,
    "HISSpecimenMaster": 62,
    "HISTestMaster": 3091,
    "MethodMaster": 49,
    "MigrationHistory": 13,
    "PatientDetails": 1820,
    "PatientVisit": 1832,
    "RadiologyRequestDetail": 160,
    "RadiologyResultDetail": 51,
    "ReferralDoctorMaster": 240,
    "RefreshTokens": 5,
    "RoleMenuPermission": 0,
    "RoleModuleMappings": 37,
    "SaleInvoice": 1654,
    "SaleInvoiceDetail": 1885,
    "SampleRejectionReasonMaster": 6,
    "SampleTypeMaster": 48,
    "TestCategoryMaster": 48,
    "TestGroupMaster": 49,
    "TestMappingMaster": 201,
    "TestParameterMappingMaster": 31,
    "TestParameters": 1,
    "TestProfileDetail": 492,
    "TestProfileMaster": 452,
    "TestRateMaster": 3270,
    "TestRequestDetails": 1574,
    "TestResultDetails": 30,
    "TestResults": 15,
    "UnitMaster": 96,
    "UserApplicationMappings": 3,
    "UserModules": 17,
}

matrix = [
    ("AspNetUsers", "Used", "Identity users / doctors", "ApplicationUser", "Account API, Auth, UI Users", "DO NOT TOUCH / CRITICAL"),
    ("AspNetRoles", "Used", "Roles", "IdentityRole", "Roles API, login claims", "DO NOT TOUCH / CRITICAL"),
    ("AspNetUserRoles", "Used", "User↔Role", "Identity", "AuthN", "DO NOT TOUCH / CRITICAL"),
    ("AspNetUserClaims", "Partially Used", "Identity claims store (0 rows)", "Identity", "Framework; unused app data", "DO NOT TOUCH"),
    ("AspNetUserLogins", "Partially Used", "External logins (0 rows)", "Identity", "Framework", "DO NOT TOUCH"),
    ("ClientApplication", "Used", "Multi-app client (DXI800)", "ClientApplication", "UserAccess, OAuth accessKey", "CRITICAL"),
    ("UserApplicationMappings", "Used", "User↔ClientApplication", "UserApplicationMapping", "Login apps list", "CRITICAL"),
    ("UserModules", "Used", "Security modules catalog", "UserModule", "Roles, QAuthorize", "CRITICAL"),
    ("RoleModuleMappings", "Used", "Module bitmask permissions", "RoleModuleMappings", "Roles Edit, QAuthorize, nav", "CRITICAL"),
    ("RoleMenuPermission", "Partially Used", "Menu overlay (0 rows; feature live)", "RoleMenuPermission", "RolesController, auth claims, Angular Role Edit", "REQUIRES BUSINESS REVIEW (feature)"),
    ("RefreshTokens", "Used", "OAuth refresh tokens", "RefreshToken", "Token endpoint", "CRITICAL"),
    ("MigrationHistory", "Used", "EF6 custom history (dbo.MigrationHistory)", "HistoryRow", "EF migrations", "DO NOT TOUCH / CRITICAL"),
    ("PatientDetails", "Used", "Patient master", "PatientDetail", "Patient API/UI, invoices, samples", "CRITICAL"),
    ("PatientVisit", "Used", "Visit management", "PatientVisit", "Visit APIs, invoice/sample FKs", "CRITICAL"),
    ("SaleInvoice", "Used", "Billing header", "SaleInvoice", "Sale Invoice UI/API", "CRITICAL"),
    ("SaleInvoiceDetail", "Used", "Billing lines", "SaleInvoiceDetail", "Sale Invoice", "CRITICAL"),
    ("TestRequestDetails", "Used", "Lab/radiology orders / samples", "TestRequestDetail", "Working board, FDD", "CRITICAL"),
    ("TestResults", "Used", "Result headers", "TestResult", "Lab result entry, approvals", "CRITICAL"),
    ("TestResultDetails", "Used", "Result parameters", "TestResultDetails", "Lab results", "CRITICAL"),
    ("TestParameters", "Partially Used", "Request-linked params (1 row)", "TestParameter", "Inflow model", "LIKELY UNUSED data; object USED"),
    ("ControlResults", "Used", "QC headers", "ControlResult", "Quality Controls UI", "DO NOT TOUCH"),
    ("ControlResultDetails", "Used", "QC details", "ControlResultDetails", "QC", "DO NOT TOUCH"),
    ("HISTestMaster", "Used", "Test catalog", "HisTestMaster", "Test Master UI", "CRITICAL"),
    ("HISParameterMaster", "Used", "Parameter catalog", "HISParameterMaster", "Parameter Master", "CRITICAL"),
    ("HISParameterRangMaster", "Used", "Reference ranges", "HISParameterRangMaster", "Ranges UI", "CRITICAL"),
    ("HISSpecimenMaster", "Used", "Specimens", "HISSpecimenMaster", "Specimen UI", "CRITICAL"),
    ("TestRateMaster", "Used", "Pricing", "TestRateMaster", "Test Rates", "CRITICAL"),
    ("TestMappingMaster", "Used", "Analyzer parameter map", "TestMappingMaster", "Mapping UI", "CRITICAL"),
    ("TestParameterMappingMaster", "Used", "Test↔parameter map", "TestParameterMappingMaster", "Mapping UI", "CRITICAL"),
    ("EquipmentMaster", "Used", "Analyzers", "EquipmentMaster", "Equipment UI", "CRITICAL"),
    ("EquipmentHeartBeat", "Partially Used", "Heartbeat (1 row)", "EquipmentHeartBeat", "Heartbeat UI", "REQUIRES BUSINESS REVIEW"),
    ("Department", "Used", "Departments", "Departments", "Setup", "CRITICAL"),
    ("UnitMaster", "Used", "Units", "UnitMaster", "Setup", "CRITICAL"),
    ("MethodMaster", "Used", "Methods", "MethodMaster", "Setup", "CRITICAL"),
    ("SampleTypeMaster", "Used", "Sample types", "SampleTypeMaster", "Setup", "DO NOT TOUCH"),
    ("ContainerMaster", "Used", "Containers", "ContainerMaster", "Setup", "DO NOT TOUCH"),
    ("ReferralDoctorMaster", "Used", "Referral doctors", "ReferralDoctorMaster", "Masters", "CRITICAL"),
    ("CorporateMaster", "Used", "Corporates", "CorporateMaster", "Masters", "CRITICAL"),
    ("TestGroupMaster", "Used", "Test groups", "TestGroupMaster", "Masters", "DO NOT TOUCH"),
    ("TestCategoryMaster", "Used", "Test categories", "TestCategoryMaster", "Masters", "DO NOT TOUCH"),
    ("TestProfileMaster", "Used", "Profiles", "TestProfileMaster", "Profiles UI", "CRITICAL"),
    ("TestProfileDetail", "Used", "Profile lines", "TestProfileDetail", "Profiles", "CRITICAL"),
    ("SampleRejectionReasonMaster", "Used", "Rejection reasons", "SampleRejectionReasonMaster", "Reject workflow", "DO NOT TOUCH"),
    ("RadiologyRequestDetail", "Used", "Radiology orders", "RadiologyRequestDetail", "Radiology modules", "CRITICAL"),
    ("RadiologyResultDetail", "Used", "Radiology results", "RadiologyResultDetail", "Radiology", "CRITICAL"),
]

rows = []
for name, clas, purpose, entity, refs, risk in matrix:
    rows.append([name, str(row_map.get(name, "?")), clas, purpose, entity, refs, risk])

add_table(
    ["Table", "Rows", "Class", "Purpose", "EF Entity", "Primary Referenced By", "Risk"],
    rows,
)

h("3.1 Objects in Code but NOT in ZoryaLMS", 2)
add_table(
    ["Object", "Evidence", "Classification", "Risk"],
    [
        [
            "TestResultsHist",
            "DTO [Table]; no DbSet; table absent",
            "Legacy / Not Referenced in live DB",
            "LIKELY UNUSED (entity only) — DO NOT DROP until confirmed never needed for audit",
        ],
        ["TestResultDetailsHist", "Same as above", "Legacy", "LIKELY UNUSED (entity only)"],
        [
            "vwTestMaster / vwTestReq / vwHISTestParameters / vw_HisParamMaster",
            "SQL_Script only; LISStaging; 0 C# refs",
            "Legacy (external/staging)",
            "REQUIRES BUSINESS REVIEW — not in ZoryaLMS",
        ],
        ["HisParamMaster", "SQL_Script create; not in ZoryaLMS", "Legacy", "REQUIRES BUSINESS REVIEW"],
        ["ExecutionLog", "SQL_Script create; not in ZoryaLMS", "Legacy", "REQUIRES BUSINESS REVIEW"],
    ],
)

h("Section 4 — Column Usage Matrix")
p(
    "Full column-by-column citation of every column across 45 tables exceeds practical report size without "
    "automated static analysis of every property access. Below: methodology + high-signal findings. "
    "Column inventory was extracted from sys.columns (hundreds of columns)."
)
p("Methodology:", bold=True)
bullets(
    [
        "EF-mapped properties on DbSet entities → treated as Referenced unless proven dead.",
        "Columns only in SQL_Script dumps with old names → Legacy relative to current EF model.",
        "Write-only / read-only distinctions require runtime tracing; marked Unknown where not proven.",
    ]
)
p("High-signal column findings:", bold=True)
add_table(
    ["Table", "Column / Theme", "Finding", "Classification", "Risk"],
    [
        [
            "PatientDetails",
            "MRNo, VisitId, Prefix",
            "Added via migrations/scripts; used in registration & uniqueness indexes",
            "Referenced",
            "DO NOT TOUCH",
        ],
        ["PatientDetails", "Address", "add-patient-address.sql + UI", "Referenced", "DO NOT TOUCH"],
        ["PatientVisit", "VisitId", "UX unique index; visit workflow", "Referenced", "CRITICAL"],
        [
            "SaleInvoice / Detail",
            "Discount/Tax/TestProfileId/nullable RequestDetailId",
            "Multiple enhancement scripts; Sale Invoice UI",
            "Referenced",
            "DO NOT TOUCH",
        ],
        [
            "Department",
            "ProcessingCategory",
            "department-processing-category.sql + API",
            "Referenced",
            "DO NOT TOUCH",
        ],
        [
            "AspNetUsers",
            "Designation, Signature fields",
            "Doctor designation/signature feature",
            "Referenced",
            "DO NOT TOUCH",
        ],
        [
            "TestRequestDetails",
            "FDD/radiology status fields",
            "fdd-sample-radiology-schema",
            "Referenced",
            "CRITICAL",
        ],
        [
            "RoleMenuPermission",
            "MenuKey + Can* flags",
            "Menu overlay; 0 rows today",
            "Referenced (schema) / Write-capable",
            "REQUIRES BUSINESS REVIEW (adoption)",
        ],
        [
            "HISParameterMaster",
            "Lab-result-restructure renames",
            "sp_rename in migrations",
            "Referenced",
            "DO NOT TOUCH",
        ],
        [
            "SQL_Script TestResults dump",
            "Legacy column names",
            "Does not match live EF model",
            "Legacy",
            "N/A (not live)",
        ],
    ],
)
p(
    "Duplicate purpose (business): PatientDetails.VisitId historically coexisted with PatientVisit table — "
    "both exist and are Used. Treat as dual-model requiring BUSINESS REVIEW for future consolidation — "
    "NOT safe to remove either."
)

h("Section 5 — Stored Procedure Matrix")
p("Live ZoryaLMS: 0 user stored procedures.")
add_table(
    ["Procedure", "Called From", "Classification", "Risk"],
    [
        ["(none in ZoryaLMS)", "—", "N/A", "—"],
        [
            "sys.sp_rename (system)",
            "EF migrations / lab-result-restructure-schema.sql only",
            "System DDL helper",
            "DO NOT TOUCH",
        ],
    ],
)
p(
    "Evidence: No CREATE PROCEDURE in solution; no CommandType.StoredProcedure / ExecuteSqlCommand calling "
    "app SPs; GenericRepository.ExecuteSQL has zero callers."
)

h("Section 6 — View Matrix")
p("Live ZoryaLMS: 0 views.")
add_table(
    ["View", "Location", "Used By", "Classification", "Risk"],
    [
        ["(none in ZoryaLMS)", "—", "—", "N/A", "—"],
        [
            "vwTestMaster",
            "SQL_Script\\vwTestMaster.sql (LISStaging)",
            "No C# refs found",
            "Legacy / Not Referenced by app",
            "REQUIRES BUSINESS REVIEW (HIS bridge?)",
        ],
        ["vwTestReq", "SQL_Script\\vwTestReq.sql", "No C# refs", "Legacy", "REQUIRES BUSINESS REVIEW"],
        [
            "vwHISTestParameters",
            "SQL_Script\\vwHISTestParameters.sql",
            "No C# refs",
            "Legacy",
            "REQUIRES BUSINESS REVIEW",
        ],
        [
            "vw_HisParamMaster",
            "SQL_Script\\vw_HisParamMaster.sql",
            "Reads HisParamMaster",
            "Legacy",
            "REQUIRES BUSINESS REVIEW",
        ],
    ],
)

h("Section 7 — Function Matrix")
p("Live ZoryaLMS: 0 user-defined functions. No SQL CLR / scalar / table-valued functions to classify.")

h("Section 8 — Index Analysis")
p(
    "DMV sys.dm_db_index_usage_stats resets on SQL Server service restart. Zero-read indexes below are "
    "CANDIDATES FOR REVIEW ONLY — not safe to drop."
)
p("Recommendations (do not implement in this engagement):", bold=True)
bullets(
    [
        "Retain all PK clustered indexes.",
        "Review unique indexes UX_PatientDetails_MRNo / VisitId and UX_PatientVisit_VisitId — business uniqueness; DO NOT TOUCH without product sign-off.",
        "Indexes with 0 reads but non-zero writes since restart (e.g. RoleModuleMappings IX_RoleId/ApplicationId, SaleInvoice FKs) — LIKELY still needed for join/FK enforcement patterns; REQUIRES BUSINESS REVIEW + longer DMV window.",
        "AspNetUserClaims.IX_UserId / AspNetUserLogins indexes — Identity framework; DO NOT TOUCH despite 0 usage.",
        "No duplicate identical indexes were proven in this pass; overlapping nonclustered indexes may exist on SaleInvoice* FK columns — optimization opportunity only.",
        "Missing-index DMV not used as authority here (can recommend harmful indexes in OLTP LIS workloads).",
    ]
)
add_table(
    ["Table", "Index", "Reads*", "Writes*", "Note"],
    [
        ["RoleModuleMappings", "IX_RoleId / IX_ApplicationId", "0", "117", "High writes; keep pending review"],
        ["SaleInvoiceDetail", "IX_TestId / IX_RequestDetailId", "0", "20", "FK support likely"],
        ["SaleInvoice", "IX_PatientId / IX_RequestDetailId", "0", "14", "FK support likely"],
        ["PatientDetails", "UX_MRNo / UX_VisitId", "0", "4", "Unique business keys"],
        ["PatientVisit", "UX_VisitId", "0", "9", "Unique business key"],
        ["RoleMenuPermission", "IX_ModuleId (+ others)", "0", "2", "New table; low traffic"],
        ["AspNetUserClaims", "IX_UserId", "0", "0", "Identity — keep"],
        ["TestProfileDetail", "IX_TestProfileId / IX_TestId", "0", "0", "Likely needed when profiles hit"],
    ],
)
p("*Since last SQL Server restart — not lifetime truth.")

h("Section 9 — Legacy Objects")
bullets(
    [
        "SQL_Script\\ folder: LISStaging / BeckmanLIS / NEOSOFT HIS bridge scripts — not applied as live ZoryaLMS objects.",
        "EF migration named fresh (Lis.Api and DataModel): bootstrap/reset-style artifact — obsolete naming; schema effects may still be foundational — DO NOT TOUCH history rows.",
        "DataModel migrations: _fixed, PatientMrVisitPrefix, FddSampleRadiology — feature evolution; keep MigrationHistory.",
        "Certification / QA scripts (SeedQa*, cert-seed-*, RoleBasedUiCertification, RemoveCrudTestData): operational tooling, not unused DB objects.",
        "Hist DTOs (TestResultsHist*): previous audit-history design never wired to DbSet in current solution.",
    ]
)

h("Section 10 — Duplicate Structures")
add_table(
    ["Theme", "Objects", "Assessment", "Action"],
    [
        [
            "Visit identity",
            "PatientDetails.VisitId vs PatientVisit",
            "Both Used; overlapping concepts",
            "BUSINESS REVIEW for consolidation",
        ],
        [
            "Parameter master naming",
            "HISParameterMaster vs HisParamMaster (script)",
            "Different eras/DBs",
            "Do not merge blindly",
        ],
        [
            "Test results",
            "TestResults vs TestResultsHist (code only)",
            "Hist unused in live DB",
            "Review audit requirements",
        ],
        [
            "Permission models",
            "RoleModuleMappings + RoleMenuPermission",
            "Intentional overlay design",
            "Keep both",
        ],
        [
            "Migration history naming",
            "MigrationHistory vs __MigrationHistory convention",
            "Custom HistoryContext maps dbo.MigrationHistory",
            "DO NOT TOUCH",
        ],
        [
            "EF dual contexts",
            "ApplicationDBContext + IdentityDbContext",
            "Same DefaultConnection; split models",
            "Architecture — keep",
        ],
    ],
)

h("Section 11 — Potential Cleanup Candidates")
p("None marked SAFE TO REMOVE. Candidates for future cleanup discussion only:", bold=True)
add_table(
    ["Candidate", "Type", "Evidence", "Suggested Future Class", "Blockers"],
    [
        [
            "TestResultsHist / TestResultDetailsHist entities",
            "Code artifacts",
            "No DbSet; tables absent",
            "LIKELY UNUSED code",
            "May be needed for future audit",
        ],
        [
            "SQL_Script views/tables",
            "Repo scripts",
            "Not in ZoryaLMS; no C# callers",
            "Legacy scripts archive",
            "External HIS may still use elsewhere",
        ],
        [
            "GenericRepository.ExecuteSQL",
            "Dead code path",
            "Zero callers",
            "Code cleanup candidate",
            "Out of DB scope",
        ],
        ["Unused using System.Data.SqlClient", "Code", "Import only", "Code cleanup", "—"],
        [
            "RoleMenuPermission empty rows",
            "Data",
            "Feature adopted, 0 overlay rows",
            "Normal empty state",
            "Not removable",
        ],
        [
            "Zero-read NC indexes",
            "Indexes",
            "DMV since restart",
            "Review after 30–90 days stats",
            "Do not drop now",
        ],
    ],
)

h("Section 12 — Objects Requiring Business Decision")
bullets(
    [
        "Whether Visit fields on PatientDetails can eventually be deprecated in favor of PatientVisit only.",
        "Whether hist tables should be implemented for compliance or code entities removed.",
        "Whether SQL_Script HIS bridge remains a supported integration path.",
        "Whether RoleMenuPermission overlay will be mandatory for all roles or remain optional.",
        "Whether EquipmentHeartBeat and Control* low-volume tables are production-critical for all sites.",
        "Index retention policy after sustained production DMV collection.",
    ]
)

h("Section 13 — Objects That Must Never Be Removed")
bullets(
    [
        "All AspNet* Identity tables",
        "ClientApplication, UserApplicationMappings, UserModules, RoleModuleMappings, RefreshTokens",
        "MigrationHistory",
        "PatientDetails, PatientVisit, SaleInvoice*, TestRequestDetails, TestResults*, HIS* masters, TestRateMaster",
        "RadiologyRequestDetail, RadiologyResultDetail",
        "EquipmentMaster, Department, referral/corporate/profile masters used by UI",
        "All Primary Keys and Foreign Keys supporting the above",
    ]
)

h("Section 14 — Recommendations")
bullets(
    [
        "Do not execute any DROP/ALTER cleanup from this report.",
        "Archive SQL_Script\\ as legacy HIS/staging documentation; confirm with integration owners before deleting from source control.",
        "Decide fate of Hist DTO classes in a code-debt sprint (separate from DB drops).",
        "Keep RoleMenuPermission; seed overlays only when product requires menu-level restriction.",
        "Collect index usage for ≥30 days of production-like load before any index drop proposals.",
        "Add a living schema ownership matrix (Table → Product Owner → Module) for future change control.",
        "Prefer retiring dead C# paths (ExecuteSQL) over touching schema.",
        "Re-run this analysis after major modules (billing, radiology, permissions) change.",
    ]
)

h("Appendix A — EF Context Map", 2)
bullets(
    [
        "ApplicationDBContext — LIS.DataModel\\DAL\\ApplicationDBContext.cs — domain DbSets",
        "IdentityDbContext — web\\Lis.Api\\Models\\IdentityModels.cs — security DbSets",
        "ApplicationHistoryContext — maps dbo.MigrationHistory",
    ]
)

h("Appendix B — Evidence Sources", 2)
bullets(
    [
        "SQL Server catalog views: sys.tables, columns, indexes, foreign_keys, key_constraints, dm_db_index_usage_stats",
        "Solution search: ExecuteSql*, SqlQuery, Dapper, EXEC, CREATE PROCEDURE/VIEW",
        "EF DbSet and [Table] attribute inventory",
        "Scripts\\ and SQL_Script\\ file inventory",
        "Row counts queried from ZoryaLMS on analysis date",
        "Peer re-verification (architect/DBA): orphan candidates confirmed ABSENT from live ZoryaLMS; CRUD-/QA row samples present",
    ]
)

h("Appendix C — Engagement Boundaries", 2)
p(
    "Primary analysis did not drop or alter any live object. This revised report appends OPTIONAL, guarded "
    "DROP/DELETE scripts for future use under change control. Scripts are not authorized for production "
    "execution without written approval. Report path: I:\\Projects\\LIS\\DBRemoval.docx"
)

h("Appendix D — Constraint & FK Summary", 2)
p(
    "45 primary keys, 42 foreign keys, 22 default constraints, 0 check constraints, 0 unique constraints "
    "(unique indexes used instead). FK graph primarily links transactions (SaleInvoice*, TestRequestDetails, "
    "Radiology*, TestResults*) to PatientDetails/PatientVisit and masters. No orphan-table FKs pointing to "
    "missing parents were detected in catalog metadata. Potential orphan *rows* were not exhaustively scanned."
)

h("Appendix E — Triggers / Sequences / UDTs / Synonyms", 2)
p(
    "Counts are all zero in live ZoryaLMS. Nothing to classify as Used/Unused for these categories. "
    "Any such objects in other environments (staging HIS) are out of scope of this database instance."
)

doc.add_page_break()

# ========== ARCHITECT / DBA PEER REVIEW ==========
h("Section 15 — Senior Architect & Senior DBA Peer Review (Errata / Amendments)")
p(
    "This section records a second-pass review of the original analysis. Goal: correct overstatements, "
    "tighten risk language, and attach only defensibly scoped SQL cleanup scripts."
)

p("15.1 Review verdict", bold=True)
bullets(
    [
        "Core conclusion STANDS: live ZoryaLMS has no unused user tables that are safe to drop. All 45 tables are EF/Identity-backed.",
        "Original report correctly refused SAFE TO REMOVE for production schema objects.",
        "Amendment: clarify that 'Partially Used' (0-row tables like AspNetUserClaims, RoleMenuPermission) means schema is required by runtime/framework/feature — empty ≠ droppable.",
        "Amendment: DMV zero-read indexes must NOT be presented as unused lifetime indexes; label as 'insufficient evidence window'.",
        "Amendment: SQL_Script orphans are repository/legacy-integration artifacts, not live ZoryaLMS catalog objects — DROP scripts against ZoryaLMS are no-ops today (verified).",
        "Amendment: data-level DELETE of CRUD-% / QA seed rows is the only cleanup with positive evidence of non-production identifiers in live DB (sample counts observed at review time).",
        "Gap remaining: column-level 'never referenced' proof is incomplete without full static property graph + runtime telemetry — keep Unknown where unproven.",
        "Gap remaining: cross-database HIS/LISStaging was not inventoried; do not assume SQL_Script objects are unused outside ZoryaLMS.",
    ]
)

p("15.2 Classification corrections", bold=True)
add_table(
    ["Item", "Original", "Amended", "Rationale"],
    [
        [
            "AspNetUserClaims / AspNetUserLogins",
            "Partially Used",
            "Used (framework-required) / empty data",
            "ASP.NET Identity schema contract; 0 rows does not authorize DROP",
        ],
        [
            "RoleMenuPermission",
            "Partially Used / business review",
            "Used (feature schema) / empty overlay data",
            "API + UI + claims path live; dropping breaks menu-permission feature",
        ],
        [
            "Zero-read NC indexes",
            "Potential cleanup",
            "Insufficient stats — DO NOT DROP",
            "DMV resets on service restart; writes observed on several",
        ],
        [
            "TestResultsHist entities",
            "LIKELY UNUSED",
            "LIKELY UNUSED code artifacts; tables absent",
            "Safe to consider code removal later; nothing to DROP in ZoryaLMS",
        ],
        [
            "SQL_Script views/tables",
            "REQUIRES BUSINESS REVIEW",
            "Unchanged + verify external DBs before any DROP elsewhere",
            "Not present in ZoryaLMS (re-verified)",
        ],
    ],
)

p("15.3 Objects still NEVER to drop from ZoryaLMS", bold=True)
bullets(
    [
        "All AspNet* tables, ClientApplication, UserModules, RoleModuleMappings, RoleMenuPermission, RefreshTokens, MigrationHistory",
        "All clinical/transaction/master tables listed in Section 3 as Used/CRITICAL",
        "All PKs/FKs on the above",
    ]
)

p("15.4 What the appended SQL is allowed to target", bold=True)
bullets(
    [
        "SCRIPT A — DROP IF EXISTS for objects that are NOT in live ZoryaLMS (defensive cleanup if someone applied SQL_Script to the wrong database).",
        "SCRIPT B — DELETE rows matching CRUD-% test identifiers (mirrors Scripts\\RemoveCrudTestData.sql).",
        "SCRIPT C — OPTIONAL commented DELETE for @zorya.test QA users (DISABLED by default — identity FKs/risk).",
        "SCRIPT D — OPTIONAL commented index drops (DISABLED — requires 30–90 day DMV proof + DBA approval).",
        "No script drops live EF-mapped business tables.",
    ]
)

doc.add_page_break()

h("Section 16 — Appended DROP / DELETE SQL Scripts (Guarded)")
p(
    "WARNING: Execute only on a restored backup copy first. Take a full backup of ZoryaLMS before any run. "
    "Obtain CAB / change-ticket approval. Scripts use transactions where deletes occur. "
    "Drop scripts use IF OBJECT_ID guards and are expected to be no-ops on current ZoryaLMS.",
    bold=True,
)

def add_sql_block(title, sql_text):
    h(title, 2)
    # Word: add as preformatted paragraphs
    for line in sql_text.strip("\n").split("\n"):
        para = doc.add_paragraph()
        run = para.add_run(line if line != "" else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(8)


add_sql_block(
    "16.1 SCRIPT A — Defensive DROP of legacy/staging objects (not present in current ZoryaLMS)",
    r"""
/*
 =============================================================================
 SCRIPT A — Defensive DROP IF EXISTS
 Target: dbo objects from legacy SQL_Script / hist design that are NOT part of
         the current EF model and were confirmed ABSENT from live ZoryaLMS at
         review time. Safe no-op if objects do not exist.
 DO NOT run against LISStaging/HIS bridge DBs without integration-owner approval.
 =============================================================================
*/
USE [ZoryaLMS];
GO
SET NOCOUNT ON;
SET XACT_ABORT ON;

PRINT 'SCRIPT A starting — defensive DROP IF EXISTS (legacy/staging names)...';

-- Views (SQL_Script legacy)
IF OBJECT_ID(N'dbo.vwTestMaster', N'V') IS NOT NULL
BEGIN
    DROP VIEW dbo.vwTestMaster;
    PRINT 'Dropped view dbo.vwTestMaster';
END
ELSE PRINT 'Skip — dbo.vwTestMaster not found';

IF OBJECT_ID(N'dbo.vwTestReq', N'V') IS NOT NULL
BEGIN
    DROP VIEW dbo.vwTestReq;
    PRINT 'Dropped view dbo.vwTestReq';
END
ELSE PRINT 'Skip — dbo.vwTestReq not found';

IF OBJECT_ID(N'dbo.vwHISTestParameters', N'V') IS NOT NULL
BEGIN
    DROP VIEW dbo.vwHISTestParameters;
    PRINT 'Dropped view dbo.vwHISTestParameters';
END
ELSE PRINT 'Skip — dbo.vwHISTestParameters not found';

IF OBJECT_ID(N'dbo.vw_HisParamMaster', N'V') IS NOT NULL
BEGIN
    DROP VIEW dbo.vw_HisParamMaster;
    PRINT 'Dropped view dbo.vw_HisParamMaster';
END
ELSE PRINT 'Skip — dbo.vw_HisParamMaster not found';

-- Legacy staging / unused hist tables (NOT in current EF DbSets; absent on review)
IF OBJECT_ID(N'dbo.HisParamMaster', N'U') IS NOT NULL
BEGIN
    DROP TABLE dbo.HisParamMaster;
    PRINT 'Dropped table dbo.HisParamMaster';
END
ELSE PRINT 'Skip — dbo.HisParamMaster not found';

IF OBJECT_ID(N'dbo.ExecutionLog', N'U') IS NOT NULL
BEGIN
    DROP TABLE dbo.ExecutionLog;
    PRINT 'Dropped table dbo.ExecutionLog';
END
ELSE PRINT 'Skip — dbo.ExecutionLog not found';

IF OBJECT_ID(N'dbo.TestResultsHist', N'U') IS NOT NULL
BEGIN
    DROP TABLE dbo.TestResultsHist;
    PRINT 'Dropped table dbo.TestResultsHist';
END
ELSE PRINT 'Skip — dbo.TestResultsHist not found';

IF OBJECT_ID(N'dbo.TestResultDetailsHist', N'U') IS NOT NULL
BEGIN
    DROP TABLE dbo.TestResultDetailsHist;
    PRINT 'Dropped table dbo.TestResultDetailsHist';
END
ELSE PRINT 'Skip — dbo.TestResultDetailsHist not found';

PRINT 'SCRIPT A complete.';
GO
""",
)

add_sql_block(
    "16.2 SCRIPT B — DELETE certification/CRUD test data only (row cleanup, not schema)",
    r"""
/*
 =============================================================================
 SCRIPT B — DELETE CRUD-% prefixed test data only
 Mirrors Scripts\RemoveCrudTestData.sql
 Does NOT delete production-shaped rows. Review rowcounts before COMMIT.
 At review time sample: ~1 CRUD patient, ~1 CRUD invoice (counts change over time).
 =============================================================================
*/
USE [ZoryaLMS];
GO
SET NOCOUNT ON;
SET XACT_ABORT ON;

BEGIN TRANSACTION;

-- Preview counts (comment out if running non-interactively after review)
SELECT 'PatientDetails' AS T, COUNT(*) AS Cnt FROM dbo.PatientDetails WHERE HisPatientId LIKE N'CRUD%';
SELECT 'SaleInvoice' AS T, COUNT(*) AS Cnt FROM dbo.SaleInvoice WHERE InvoiceNo LIKE N'CRUD%';
SELECT 'TestRequestDetails' AS T, COUNT(*) AS Cnt FROM dbo.TestRequestDetails WHERE SampleNo LIKE N'CRUD%';
SELECT 'HISTestMaster' AS T, COUNT(*) AS Cnt FROM dbo.HISTestMaster WHERE HISTestCode LIKE N'CRUD%';

DELETE d FROM dbo.TestResultDetails d
INNER JOIN dbo.TestResults r ON r.Id = d.TestResultId
WHERE r.SampleNo LIKE N'CRUD%';

DELETE FROM dbo.TestResults WHERE SampleNo LIKE N'CRUD%';

DELETE d FROM dbo.ControlResultDetails d
INNER JOIN dbo.ControlResults c ON c.Id = d.ControlResultId
WHERE c.SampleNo LIKE N'CRUD%';

DELETE FROM dbo.ControlResults WHERE SampleNo LIKE N'CRUD%';

DELETE FROM dbo.TestParameters WHERE HISTestCode LIKE N'CRUD%';

DELETE FROM dbo.SaleInvoiceDetail
WHERE SaleInvoiceId IN (SELECT Id FROM dbo.SaleInvoice WHERE InvoiceNo LIKE N'CRUD%');

DELETE FROM dbo.SaleInvoice WHERE InvoiceNo LIKE N'CRUD%';

DELETE FROM dbo.TestRequestDetails WHERE SampleNo LIKE N'CRUD%';

-- Visits linked only to CRUD patients (if VisitId/MR pattern used)
DELETE FROM dbo.PatientVisit
WHERE PatientId IN (SELECT Id FROM dbo.PatientDetails WHERE HisPatientId LIKE N'CRUD%');

DELETE FROM dbo.PatientDetails WHERE HisPatientId LIKE N'CRUD%';

DELETE FROM dbo.TestRateMaster WHERE TestId IN (
    SELECT Id FROM dbo.HISTestMaster WHERE HISTestCode LIKE N'CRUD%'
);

DELETE d FROM dbo.TestProfileDetail d
INNER JOIN dbo.TestProfileMaster p ON p.Id = d.TestProfileId
WHERE p.Code LIKE N'CRUD%';

DELETE FROM dbo.TestProfileMaster WHERE Code LIKE N'CRUD%';
DELETE FROM dbo.TestMappingMaster WHERE HISTestCode LIKE N'CRUD%';

DELETE FROM dbo.HISParameterRangMaster WHERE HisParameterId IN (
    SELECT Id FROM dbo.HISParameterMaster WHERE HISTestCode LIKE N'CRUD%'
);
DELETE FROM dbo.HISParameterMaster WHERE HISTestCode LIKE N'CRUD%';
DELETE FROM dbo.HISTestMaster WHERE HISTestCode LIKE N'CRUD%';

DELETE FROM dbo.EquipmentHeartBeat WHERE AccessKey LIKE N'CRUD%';
DELETE FROM dbo.EquipmentMaster WHERE AccessKey LIKE N'CRUD%';
DELETE FROM dbo.CorporateMaster WHERE Code LIKE N'CRUD%';
DELETE FROM dbo.ReferralDoctorMaster WHERE Code LIKE N'CRUD%';
DELETE FROM dbo.TestCategoryMaster WHERE Code LIKE N'CRUD%';
DELETE FROM dbo.TestGroupMaster WHERE Code LIKE N'CRUD%';
DELETE FROM dbo.ContainerMaster WHERE Code LIKE N'CRUD%';
DELETE FROM dbo.SampleTypeMaster WHERE Code LIKE N'CRUD%';
DELETE FROM dbo.MethodMaster WHERE Code LIKE N'CRUD%';
DELETE FROM dbo.UnitMaster WHERE Code LIKE N'CRUD%';
DELETE FROM dbo.HISSpecimenMaster WHERE Code LIKE N'CRUD%';
DELETE FROM dbo.Department WHERE Code LIKE N'CRUD%';

-- CHANGE COMMIT to ROLLBACK for dry-run
COMMIT TRANSACTION;
-- ROLLBACK TRANSACTION;

PRINT 'SCRIPT B complete — CRUD-% test data deleted (if any).';
GO
""",
)

add_sql_block(
    "16.3 SCRIPT C — OPTIONAL QA identity users (DISABLED — uncomment only with IAM approval)",
    r"""
/*
 =============================================================================
 SCRIPT C — OPTIONAL delete of @zorya.test QA users
 STATUS: DISABLED BY DEFAULT (entire batch commented)
 RISK: HIGH — may break AspNetUserRoles / mappings / audit trails
 Review-time sample: ~5 emails like %@zorya.test
 =============================================================================
*/
USE [ZoryaLMS];
GO
/*
SET NOCOUNT ON;
SET XACT_ABORT ON;
BEGIN TRANSACTION;

DECLARE @qa TABLE (UserId NVARCHAR(128) PRIMARY KEY);
INSERT INTO @qa (UserId)
SELECT Id FROM dbo.AspNetUsers
WHERE Email LIKE N'%@zorya.test';

SELECT u.Email, u.UserName
FROM dbo.AspNetUsers u
INNER JOIN @qa q ON q.UserId = u.Id;

DELETE FROM dbo.AspNetUserRoles WHERE UserId IN (SELECT UserId FROM @qa);
DELETE FROM dbo.AspNetUserClaims WHERE UserId IN (SELECT UserId FROM @qa);
DELETE FROM dbo.AspNetUserLogins WHERE UserId IN (SELECT UserId FROM @qa);
DELETE FROM dbo.UserApplicationMappings WHERE UserId IN (SELECT UserId FROM @qa);
DELETE FROM dbo.AspNetUsers WHERE Id IN (SELECT UserId FROM @qa);

-- COMMIT TRANSACTION;
ROLLBACK TRANSACTION;
PRINT 'SCRIPT C rolled back / disabled template.';
*/
GO
""",
)

add_sql_block(
    "16.4 SCRIPT D — OPTIONAL index drops (DISABLED — insufficient DMV evidence)",
    r"""
/*
 =============================================================================
 SCRIPT D — OPTIONAL nonclustered index drops
 STATUS: DISABLED BY DEFAULT
 Prerequisite: >=30 days production DMV with zero seeks/scans/lookups AND
               confirmed unused by query plans AND DBA + app-owner sign-off.
 DO NOT drop unique business indexes (UX_PatientDetails_*, UX_PatientVisit_*).
 DO NOT drop Identity indexes.
 =============================================================================
*/
USE [ZoryaLMS];
GO
/*
-- Example template only (names must be re-validated before use):
-- IF EXISTS (SELECT 1 FROM sys.indexes WHERE name = N'IX_Example' AND object_id = OBJECT_ID(N'dbo.SomeTable'))
--     DROP INDEX IX_Example ON dbo.SomeTable;
PRINT 'SCRIPT D is intentionally empty/disabled.';
*/
GO
""",
)

add_sql_block(
    "16.5 SCRIPT E — Explicitly FORBIDDEN drops (documentation — do not execute)",
    r"""
/*
 =============================================================================
 SCRIPT E — FORBIDDEN — listed to prevent accidental generation/use
 The following must NEVER be dropped by cleanup automation for ZoryaLMS:
   AspNetUsers, AspNetRoles, AspNetUserRoles, AspNetUserClaims, AspNetUserLogins
   ClientApplication, UserApplicationMappings, UserModules, RoleModuleMappings
   RoleMenuPermission, RefreshTokens, MigrationHistory
   PatientDetails, PatientVisit, SaleInvoice, SaleInvoiceDetail
   TestRequestDetails, TestResults, TestResultDetails, TestParameters
   HISTestMaster, HISParameterMaster, HISParameterRangMaster, HISSpecimenMaster
   TestRateMaster, EquipmentMaster, Department, Radiology*, TestProfile*, masters...
 If a tool proposes DROP TABLE on any of the above → REJECT.
 =============================================================================
*/
""",
)

h("Section 17 — Execution Checklist (DBA)", 2)
bullets(
    [
        "1. Full backup of ZoryaLMS (+ log backup if FULL recovery).",
        "2. Restore backup to a scratch database; run SCRIPT A/B there first.",
        "3. Compare rowcounts before/after; validate app smoke login + role edit + patient search.",
        "4. Change ticket + dual control for production.",
        "5. Prefer SCRIPT B (data) over any schema drops.",
        "6. Leave SCRIPT C/D commented unless explicitly approved.",
        "7. Never enable SCRIPT E targets.",
    ]
)

p(
    "End of revised report. Classification policy unchanged: zero SAFE TO REMOVE schema objects in live ZoryaLMS; "
    "appended SQL is guarded future-ops material only."
)

out_path = r"I:\Projects\LIS\DBRemoval.docx"
os.makedirs(os.path.dirname(out_path), exist_ok=True)
doc.save(out_path)
print("SAVED", out_path, "SIZE", os.path.getsize(out_path))
