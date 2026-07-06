# -*- coding: utf-8 -*-
"""Generate Radiology_Architecture_Analysis.docx for ZoryaLIS."""
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

OUT = r"I:\Projects\LIS\Radiology_Architecture_Analysis.docx"

def h(doc, text, level=1):
    doc.add_heading(text, level=level)

def p(doc, text, bold=False):
    para = doc.add_paragraph()
    run = para.add_run(text)
    if bold:
        run.bold = True
    return para

def bullet(doc, text):
    doc.add_paragraph(text, style='List Bullet')

def table(doc, headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = 'Table Grid'
    for i, htext in enumerate(headers):
        t.rows[0].cells[i].text = htext
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            t.rows[ri + 1].cells[ci].text = str(val)
    doc.add_paragraph()

def main():
    doc = Document()
    title = doc.add_heading('ZoryaLIS Radiology Module', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph('Architecture, Business Flow & Technical Analysis')
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph('Document Type: Reverse Engineering & Architecture Review')
    doc.add_paragraph('System: ZoryaLIS / AVILIS')
    doc.add_paragraph('Date: June 2026')
    doc.add_paragraph('Status: Analysis Only — No Code Changes')
    doc.add_page_break()

    # 1 Executive Summary
    h(doc, '1. Executive Summary', 1)
    p(doc, (
        'ZoryaLIS does not implement Radiology as a separate product module with its own Test Master. '
        'Radiology tests are ordinary rows in HISTestMaster, classified at runtime by Department assignment '
        '(department name containing "RADIOLOGY", or department code containing "RAD"). '
        'When invoiced, radiology lines create RadiologyRequestDetail records and bypass the laboratory '
        'sample workflow (TestRequestDetails / Sample Collection / Receiving). '
        'Reporting is handled through Radiology Report Entry (queue, draft, review, authorize, release) '
        'with findings stored in RadiologyResultDetail. '
        'Blood tests and radiology tests share billing (SaleInvoice / SaleInvoiceDetail) but diverge '
        'into parallel downstream pipelines after save.'
    ))

    # 2 Architecture Overview
    h(doc, '2. Architecture Overview', 1)
    h(doc, '2.1 Solution Projects Involved', 2)
    table(doc, ['Project', 'Role in Radiology'], [
        ('LIS.DtoModel', 'Entities: RadiologyRequestDetail, RadiologyResultDetail, HisTestMaster, DTOs, enums'),
        ('LIS.DataModel (LIS.DataAccess)', 'EF ApplicationDBContext, DbSets, migrations'),
        ('LIS.Businesslogic', 'SaleInvoiceManager, RadiologyReportManager, ReportManager'),
        ('web/Lis.Api', 'RadiologyReportController, SaleInvoiceController, OperationalReportsController'),
        ('web/Lis.Web', 'Angular: radiology-report-entry, FddReportComponent (radiology reports), sale-invoice-form'),
        ('Scripts/fdd-sample-radiology-schema.sql', 'DDL for radiology tables + module seed'),
    ])

    h(doc, '2.2 Layered Components', 2)
    table(doc, ['Layer', 'Artifacts'], [
        ('Controllers', 'RadiologyReportController, SaleInvoiceController, OperationalReportsController'),
        ('Managers', 'RadiologyReportManager (IRadiologyReportManager), SaleInvoiceManager (radiology branch), ReportManager (radiology reports)'),
        ('Repositories', 'ModuleRepo<RadiologyRequestDetail>, ModuleRepo<RadiologyResultDetail> (GenericUnitOfWork pattern)'),
        ('DTOs / Entities', 'RadiologyRequestDetail, RadiologyResultDetail, BillableItemLookup, SampleWorkflow DTOs'),
        ('Angular Services', 'SampleWorkflowService (radiology APIs), MasterService (invoice save, billable lookup)'),
        ('Angular Components', 'RadiologyReportEntryComponent, FddReportComponent, SaleInvoiceFormComponent'),
        ('SQL', 'fdd-sample-radiology-schema.sql; EF migration 202606191317350_FddSampleRadiology'),
        ('Stored Procedures', 'None dedicated to radiology workflow; legacy vwTestReq.sql excludes RADIOLOGY dept from lab import'),
    ])

    h(doc, '2.3 Architecture Diagram (Text)', 2)
    p(doc, (
        '┌─────────────────────────────────────────────────────────────────────────┐\n'
        '│                         ANGULAR PORTAL (Lis.Web)                        │\n'
        '│  Sale Invoice Form ──► MasterService ──► api/SaleInvoice              │\n'
        '│  Radiology Report Entry ──► SampleWorkflowService ──► api/RadiologyReport│\n'
        '│  Radiology Reports ──► FddReportComponent ──► api/Reports/*Radiology*   │\n'
        '└───────────────────────────────┬─────────────────────────────────────────┘\n'
        '                                │ HTTP (Web API 2)\n'
        '┌───────────────────────────────▼─────────────────────────────────────────┐\n'
        '│                    Lis.Api Controllers + QAuthorize                      │\n'
        '└───────────────────────────────┬─────────────────────────────────────────┘\n'
        '                                │\n'
        '┌───────────────────────────────▼─────────────────────────────────────────┐\n'
        '│              LIS.Businesslogic Managers (Scoped via SimpleInjector)      │\n'
        '│   SaleInvoiceManager          RadiologyReportManager      ReportManager    │\n'
        '└───────────────────────────────┬─────────────────────────────────────────┘\n'
        '                                │ ModuleRepo<T> + GenericUnitOfWork\n'
        '┌───────────────────────────────▼─────────────────────────────────────────┐\n'
        '│         SQL Server (AVSLIS) — Entity Framework 6 ApplicationDBContext      │\n'
        '│  HISTestMaster, Department, SaleInvoice, SaleInvoiceDetail,              │\n'
        '│  TestRequestDetails (lab), RadiologyRequestDetail, RadiologyResultDetail   │\n'
        '└─────────────────────────────────────────────────────────────────────────┘'
    ))

    doc.add_page_break()

    # Phase 2 Business Q&A
    h(doc, '3. Business Understanding (Phase 2)', 1)

    h(doc, 'Q1 — Where is a Radiology Test created?', 2)
    bullet(doc, 'Master screen: Test Master UI at routes /test-master, /test-master/create, /test-master/:id (HisTestMaster CRUD via api/HisTest)')
    bullet(doc, 'Database table: HISTestMaster (not a separate radiology master table)')
    bullet(doc, 'Classification: Assign DepartmentCode to a department whose Name contains "RADIOLOGY" (e.g. DEP00011 / RADIOLOGY)')
    bullet(doc, 'API: POST/PUT api/HisTest for master maintenance; radiology workflow request created at invoice save or POST api/RadiologyReport')
    bullet(doc, 'Manager: HISTestMasterManager for master; SaleInvoiceManager.LinkRadiologyRequestsToLines or RadiologyReportManager.CreateRequest for requests')
    bullet(doc, 'Angular: TestCreateComponent / TestEditComponent for master; Sale Invoice form for billing-triggered requests; Radiology Report Entry for manual CreateRequest (API exists, UI primarily uses invoice path)')

    h(doc, 'Q2 — Separate Test Master vs HIS Test Master?', 2)
    p(doc, 'Radiology REUSES HisTestMaster (table HISTestMaster). Evidence: HisTestMaster.cs has no TestType/Radiology flag; SaleInvoiceManager and GetBillableItems load tests from testRepo (HisTestMaster); radiology detection is department-based only.')

    h(doc, 'Q3 — How Blood vs Radiology is distinguished', 2)
    table(doc, ['Mechanism', 'Implementation'], [
        ('Dedicated TestType enum', 'NOT used for radiology'),
        ('Database flag on HISTestMaster', 'NONE'),
        ('Department', 'PRIMARY — Department.Name contains "RADIOLOGY" (case-insensitive)'),
        ('Department code fallback', 'DepartmentCode contains "RAD"'),
        ('UI LineType', 'SaleInvoiceDetail.LineType = "radiology" (NotMapped, from UI only)'),
        ('Billable lookup', 'GetBillableItems sets LineType via IsRadiologyTest()'),
        ('Code location', 'SaleInvoiceManager.IsRadiologyTest(), IsRadiologyLine()'),
    ])

    h(doc, 'Q4 — Same Test Master for Blood, Radiology, Profiles?', 2)
    p(doc, (
        'YES. HISTestMaster holds all individual tests (blood and radiology). '
        'TestProfileMaster + TestProfileDetail group multiple TestIds into packages. '
        'Profiles reference HisTestMaster.TestId only; no enforcement that profile members are blood-only. '
        'Department on each test determines routing at invoice save for individual lines; profile expansion always creates TestRequestDetails (see Phase 8 gap).'
    ))

    doc.add_page_break()

    # Phase 3 Sale Invoice
    h(doc, '4. Sale Invoice Save Flow (Phase 3)', 1)
    h(doc, '4.1 Execution Sequence', 2)
    p(doc, (
        'UI: SaleInvoiceFormComponent.onSubmit() → MasterService.saveInvoice(dto)\n'
        'API: POST api/SaleInvoice/ → SaleInvoiceController.Post(dto) → SaleInvoiceManager.Save(dto)\n'
        'Transaction: unitOfWork.BeginTransaction() → SaveCore() → Commit (rollback on failure)\n\n'
        'SaveCore order:\n'
        '1. Validate header, patient, lines; NormalizeProfileLines()\n'
        '2. Recalculate amounts; persist SaleInvoice header (invoiceRepo.Add/Update)\n'
        '3. PersistInvoiceDetails() → SaleInvoiceDetail rows (RequestDetailId cleared)\n'
        '4. LinkTestRequestsToLines() → TestRequestDetails for blood tests & profiles\n'
        '5. LinkRadiologyRequestsToLines() → RadiologyRequestDetail for radiology tests\n'
        '6. UpdateDetailRequestLinks() → writes RequestDetailId FK on lines (lab only)\n'
        '7. Optionally set header.RequestDetailId from first lab line'
    ))

    h(doc, '4.2 Decision Logic', 2)
    p(doc, 'Pseudocode reflecting actual SaleInvoiceManager implementation:')
    p(doc, (
        'foreach line in invoice lines:\n'
        '  if line has TestProfileId:\n'
        '    ExpandProfileTestRequests() → TestRequestDetail per profile test (NO radiology skip)\n'
        '    continue\n'
        '  test = HisTestMaster by line.TestId\n'
        '  if IsRadiologyLine(line, test):  // LineType=="radiology" OR IsRadiologyTest(test)\n'
        '    skip TestRequestDetails in LinkTestRequestsToLines\n'
        '  else:\n'
        '    create/find TestRequestDetail; line.RequestDetailId = request.Id\n\n'
        'foreach line (non-profile, TestId > 0):\n'
        '  if IsRadiologyLine(line, test):\n'
        '    create RadiologyRequestDetail if not exists (PatientId + HISRequestNo + HISTestCode)\n'
        '    // NO FK from SaleInvoiceDetail to RadiologyRequestDetail'
    ))

    h(doc, '4.3 Per Selection Type', 2)
    table(doc, ['User Selection', 'SaleInvoiceDetail', 'TestRequestDetails', 'RadiologyRequestDetail'], [
        ('Blood test', 'TestId set; TestProfileId null; LineType blood', 'Created/linked; RequestDetailId set', 'Not created'),
        ('Profile', 'TestProfileId set; TestId = first profile test', 'One per profile detail test', 'Not created (even if profile contains radiology tests — gap)'),
        ('Radiology test', 'TestId set; LineType radiology', 'Skipped', 'Created Pending; AccessionNo = InvoiceNo-TestCode'),
    ])

    doc.add_page_break()

    # Phase 4 Database
    h(doc, '5. Database Flow & ER Model (Phase 4)', 1)
    p(doc, 'Note: There is NO OrderDetails table. SaleInvoiceDetail is the order-line entity.')

    h(doc, '5.1 Tables Affected on Radiology Invoice Save', 2)
    table(doc, ['Table', 'PK', 'FK / Links', 'Data Inserted', 'Business Role'], [
        ('SaleInvoice', 'Id', 'PatientId, optional RequestDetailId→TestRequestDetails', 'Invoice header', 'Billing document'),
        ('SaleInvoiceDetail', 'Id', 'SaleInvoiceId, TestId→HISTestMaster, RequestDetailId→TestRequestDetails (nullable)', 'Line items', 'Billable order line; lab FK only'),
        ('HISTestMaster', 'Id', 'DepartmentCode→Department', 'No insert on save', 'Test definition'),
        ('RadiologyRequestDetail', 'Id', 'PatientId→PatientDetails', 'New radiology order', 'Radiology queue header'),
        ('RadiologyResultDetail', 'Id', 'RadiologyRequestId (unique)', 'Not on invoice save', 'Report body on entry/authorize'),
        ('TestRequestDetails', 'Id', 'PatientId', 'Not for radiology lines', 'Lab workflow'),
        ('PatientDetails', 'Id', '—', 'No insert', 'Patient context'),
    ])

    h(doc, '5.2 ER Diagram (Text)', 2)
    p(doc, (
        'PatientDetails (1) ──< (N) SaleInvoice\n'
        'SaleInvoice (1) ──< (N) SaleInvoiceDetail\n'
        'HISTestMaster (1) ──< (N) SaleInvoiceDetail\n'
        'SaleInvoiceDetail (N) ──> (0..1) TestRequestDetails  [lab lines only]\n'
        'PatientDetails (1) ──< (N) TestRequestDetails\n'
        'PatientDetails (1) ──< (N) RadiologyRequestDetail\n'
        'RadiologyRequestDetail (1) ──< (1) RadiologyResultDetail\n'
        'TestProfileMaster (1) ──< (N) TestProfileDetail ──> (N) HISTestMaster\n'
        'HISTestMaster (N) ──> (1) Department [via DepartmentCode]\n\n'
        'MISSING LINK: No FK from SaleInvoiceDetail or SaleInvoice to RadiologyRequestDetail'
    ))

    doc.add_page_break()

    # Phase 5 Blood workflow
    h(doc, '6. Blood Test Workflow (Phase 5)', 1)
    p(doc, (
        'Patient (PatientDetails) → Sale Invoice → SaleInvoiceDetail → TestRequestDetails '
        '(created on save with SampleNo, specimen, ReportStatus=New) → Sample Collection '
        '(SampleCollectionManager updates CollectedBy, dates) → Sample Receiving '
        '(SampleReceivingManager) → Recent Sample / queues → Edit Results (TestResultEditManager) '
        '→ Technician Approval → Doctor Approval → Diagnostic Report (TestReportManager / api/Reports/TestReport).\n\n'
        'Tables: PatientDetails, SaleInvoice, SaleInvoiceDetail, TestRequestDetails, TestParameters, '
        'TestResults, TestResultDetails, HISParameterMaster.'
    ))

    # Phase 6 Radiology workflow
    h(doc, '7. Radiology Workflow (Phase 6)', 1)
    table(doc, ['Step', 'UI', 'API', 'Tables'], [
        ('Patient', 'Patient Master / Registration', 'api/PatientMaster', 'PatientDetails'),
        ('Sale Invoice', 'sale-invoice-form', 'POST api/SaleInvoice', 'SaleInvoice, SaleInvoiceDetail'),
        ('Radiology Request', '(automatic on save)', 'SaleInvoiceManager', 'RadiologyRequestDetail'),
        ('Queue', 'radiology-report-entry', 'GET api/RadiologyReport/PendingQueue', 'RadiologyRequestDetail'),
        ('Report Entry', 'radiology-report-entry (modal)', 'POST api/RadiologyReport/Save', 'RadiologyResultDetail'),
        ('Authorization', 'same screen', 'POST api/RadiologyReport/Authorize', 'RadiologyResultDetail + status'),
        ('Release', 'Authorize with release=true', 'POST api/RadiologyReport/Authorize', 'ReportStatus=Released'),
        ('Reports', 'reports/radiology/*', 'api/Reports/PendingRadiology, AuthorizedRadiology, etc.', 'Read-only aggregates'),
    ])

    h(doc, '7.1 RadiologyReportStatus Lifecycle', 2)
    p(doc, 'Pending(0) → Draft(1) / UnderReview(2) → Authorized(3) → Released(4)')

    doc.add_page_break()

    # Phase 7 Separation
    h(doc, '8. Separation Logic (Phase 7) — CRITICAL', 1)
    p(doc, 'Exact code path in SaleInvoiceManager:')
    p(doc, (
        'IsRadiologyLine(line, test):\n'
        '  if line.LineType equals "radiology" (ignore case) → true\n'
        '  else return IsRadiologyTest(test)\n\n'
        'IsRadiologyTest(test):\n'
        '  lookup Department where Code = test.DepartmentCode\n'
        '  if dept.Name contains "RADIOLOGY" → true\n'
        '  if test.DepartmentCode contains "RAD" → true\n'
        '  else false\n\n'
        'LinkTestRequestsToLines: if IsRadiologyLine → continue (skip)\n'
        'LinkRadiologyRequestsToLines: if NOT IsRadiologyLine → continue (skip)'
    ))
    p(doc, 'Radiology lines do NOT populate SaleInvoiceDetail.RequestDetailId. That FK points only to TestRequestDetails.')

    # Phase 8 Profiles
    h(doc, '9. Test Profile Behaviour (Phase 8)', 1)
    p(doc, (
        'TestProfileMaster contains TestProfileDetail rows (TestId + Quantity). '
        'Validation (TestProfileMasterManager.ValidateProfile) requires at least one test, unique TestIds, positive quantity — '
        'but does NOT restrict department or radiology membership.\n\n'
        'On invoice save, ExpandProfileTestRequests creates TestRequestDetail for EVERY profile member without calling IsRadiologyTest. '
        'Therefore profiles CAN contain mixed or radiology-only tests in master data, but expansion always follows the LAB path. '
        'This is inconsistent with individual radiology line handling and is documented as a business/technical gap.'
    ))

    doc.add_page_break()

    # Phase 9 Schema
    h(doc, '10. Radiology Table Schema (Phase 9)', 1)

    h(doc, '10.1 RadiologyRequestDetail', 2)
    table(doc, ['Column', 'Type', 'Nullable', 'FK/Purpose'], [
        ('Id', 'BIGINT IDENTITY', 'No', 'PK'),
        ('PatientId', 'BIGINT', 'No', 'FK → PatientDetails'),
        ('HISRequestNo', 'NVARCHAR(20)', 'Yes', 'Usually invoice number'),
        ('AccessionNo', 'NVARCHAR(30)', 'Yes', 'InvoiceNo-TestCode on invoice path'),
        ('Modality', 'NVARCHAR(30)', 'Yes', 'From specimen name/code/description'),
        ('HISTestCode', 'NVARCHAR(20)', 'Yes', 'From HisTestMaster'),
        ('HISTestName', 'NVARCHAR(100)', 'Yes', 'Test description'),
        ('Department', 'NVARCHAR(80)', 'Yes', 'Resolved department name'),
        ('ReportStatus', 'INT', 'No', 'RadiologyReportStatus enum'),
        ('CreatedBy/On, ModifiedBy/On', 'NVARCHAR/DATETIME', 'Yes/No', 'Audit'),
    ])
    p(doc, 'Indexes: IX_RadiologyRequest_Patient, IX_RadiologyRequest_Status')

    h(doc, '10.2 RadiologyResultDetail', 2)
    table(doc, ['Column', 'Type', 'Nullable', 'FK/Purpose'], [
        ('Id', 'BIGINT IDENTITY', 'No', 'PK'),
        ('RadiologyRequestId', 'BIGINT', 'No', 'FK → RadiologyRequestDetail (UNIQUE)'),
        ('ClinicalHistory', 'NVARCHAR(MAX)', 'Yes', 'Clinical context'),
        ('Findings', 'NVARCHAR(MAX)', 'Yes', 'Mandatory for save'),
        ('Impression', 'NVARCHAR(MAX)', 'Yes', 'Mandatory for save'),
        ('Recommendation', 'NVARCHAR(MAX)', 'Yes', 'Optional'),
        ('AuthorizedBy/On', 'NVARCHAR/DATETIME', 'Yes', 'Sign-off'),
        ('DigitalSignature', 'NVARCHAR(200)', 'Yes', 'Required for authorize'),
        ('Created/Modified audit', '', '', ''),
    ])

    doc.add_page_break()

    # Phase 10 APIs
    h(doc, '11. API Documentation (Phase 10)', 1)
    table(doc, ['Route', 'Method', 'Purpose', 'Permission Module'], [
        ('api/SaleInvoice/', 'POST', 'Save invoice; triggers radiology requests', 'SaleInvoices CanAdd/Edit'),
        ('api/SaleInvoice/BillableItems', 'GET', 'Search blood/profile/radiology billable items', 'Authenticated'),
        ('api/RadiologyReport/PendingQueue', 'GET', 'Radiology work queue', 'RadiologyReportEntry CanView'),
        ('api/RadiologyReport/{id}', 'GET', 'Report detail for entry', 'RadiologyReportEntry CanView'),
        ('api/RadiologyReport', 'POST', 'Manual create radiology request', 'RadiologyReportEntry CanAdd'),
        ('api/RadiologyReport/Save', 'POST', 'Save draft / submit review', 'RadiologyReportEntry CanEdit'),
        ('api/RadiologyReport/Authorize', 'POST', 'Authorize / release', 'RadiologyReportEntry CanAuthorize'),
        ('api/Reports/PendingRadiology', 'GET', 'Pending cases report', 'RadiologyReports CanView'),
        ('api/Reports/AuthorizedRadiology', 'GET', 'Authorized reports', 'RadiologyReports CanView'),
        ('api/Reports/ModalityStatistics', 'GET', 'Modality stats', 'RadiologyReports CanView'),
        ('api/Reports/RadiologistProductivity', 'GET', 'Productivity', 'RadiologyReports CanView'),
        ('api/HisTest/*', 'CRUD', 'Test master (includes radiology tests)', 'HisTest module'),
    ])

    # Phase 11 UI
    h(doc, '12. UI Documentation (Phase 11)', 1)
    table(doc, ['Route', 'Component', 'Module Permission'], [
        ('/test-master', 'TestListComponent', 'Test master'),
        ('/test-master/create', 'TestCreateComponent', 'Create radiology-capable test'),
        ('/sale-invoices/create', 'SaleInvoiceFormComponent', 'Bill radiology via unified selector'),
        ('/radiology-report-entry', 'RadiologyReportEntryComponent', 'RadiologyReportEntry'),
        ('/reports/radiology/pending', 'FddReportComponent', 'RadiologyReports'),
        ('/reports/radiology/authorized', 'FddReportComponent', 'RadiologyReports'),
        ('/reports/radiology/modality-stats', 'FddReportComponent', 'RadiologyReports'),
        ('/reports/radiology/productivity', 'FddReportComponent', 'RadiologyReports'),
    ])
    p(doc, 'Navigation: Left nav → Operations → Radiology Report Entry; Reports → Pending/Authorized Radiology sections.')

    doc.add_page_break()

    # Phase 12 Sequence diagrams
    h(doc, '13. Sequence Diagrams (Phase 12)', 1)
    h(doc, '13.1 Radiology Test (from Invoice)', 2)
    p(doc, (
        'User → SaleInvoiceForm → POST /api/SaleInvoice\n'
        '→ SaleInvoiceManager.Save → PersistInvoiceDetails\n'
        '→ LinkRadiologyRequestsToLines → ModuleRepo.Add(RadiologyRequestDetail)\n'
        '→ SQL INSERT RadiologyRequestDetail\n'
        '← invoice id\n'
        'Later: User → RadiologyReportEntry → GET PendingQueue\n'
        '→ GET /api/RadiologyReport/{id} → POST Save → POST Authorize\n'
        '→ RadiologyResultDetail INSERT/UPDATE'
    ))

    h(doc, '13.2 Blood Test (from Invoice)', 2)
    p(doc, (
        'User → SaleInvoiceForm → POST /api/SaleInvoice\n'
        '→ LinkTestRequestsToLines → TestRequestDetail INSERT\n'
        '→ UpdateDetailRequestLinks → SaleInvoiceDetail.RequestDetailId\n'
        '→ Sample Collection / Receiving APIs on TestRequestDetails'
    ))

    h(doc, '13.3 Test Profile', 2)
    p(doc, (
        'User selects profile line → POST /api/SaleInvoice\n'
        '→ ExpandProfileTestRequests → multiple TestRequestDetail\n'
        '→ line.RequestDetailId = first request id\n'
        '(No radiology branch in profile expansion)'
    ))

    # Phase 13 Business rules
    h(doc, '14. Business Rules (Phase 13)', 1)
    rules = [
        'Radiology tests are HisTestMaster rows assigned to a RADIOLOGY department (or RAD code).',
        'Radiology invoice lines do not create TestRequestDetails.',
        'Radiology invoice lines do not link SaleInvoiceDetail.RequestDetailId.',
        'Radiology requests link to patient + HISRequestNo (invoice no) + HISTestCode for deduplication.',
        'Radiology does not use sample collection or specimen receiving workflows.',
        'Findings and Impression are mandatory before save; digital signature mandatory before authorize.',
        'Authorized/Released reports cannot be edited.',
        'Billable items require active TestRate effective on invoice date.',
        'Profile lines bill as package rate; expand to lab requests per TestProfileDetail.',
        'Lab diagnostic report uses TestReportManager; radiology uses separate report tables.',
    ]
    for r in rules:
        bullet(doc, r)

    # Phase 14 Risks
    h(doc, '15. Risks & Observations (Phase 14)', 1)
    h(doc, 'Strengths', 2)
    bullet(doc, 'Clear separation of lab vs radiology downstream tables')
    bullet(doc, 'Reuses single test master and billing model')
    bullet(doc, 'Transactional invoice save with explicit manager methods')
    bullet(doc, 'Dedicated radiology reporting lifecycle with authorization')

    h(doc, 'Weaknesses / Technical Debt', 2)
    bullet(doc, 'No FK from invoice lines to RadiologyRequestDetail — hard to trace invoice→radiology request')
    bullet(doc, 'Profile expansion ignores radiology classification')
    bullet(doc, 'IsRadiologyTest heuristic (name contains RADIOLOGY / code contains RAD) is fragile')
    bullet(doc, 'Modality derived from specimen fields — semantic misuse of HISSpecimenName')
    bullet(doc, 'No stored procedure layer; some report queries load full tables client-side in managers')
    bullet(doc, 'LineType on SaleInvoiceDetail is NotMapped — server re-derives from department if UI omits it')

    h(doc, 'Future Extensibility', 2)
    bullet(doc, 'Add explicit TestServiceType enum on HISTestMaster for reliable routing')
    bullet(doc, 'Link SaleInvoiceDetail to RadiologyRequestDetail optional FK')
    bullet(doc, 'Extend ExpandProfileTestRequests with IsRadiologyTest branching')
    bullet(doc, 'Dedicated Modality master instead of specimen field mapping')

    # Recommendations
    h(doc, '16. Recommendations (Documentation Only)', 1)
    bullet(doc, 'Document department setup: assign radiology tests only to RADIOLOGY department')
    bullet(doc, 'Avoid radiology tests inside profiles until expansion logic is aligned')
    bullet(doc, 'Use operational reports to reconcile invoice radiology lines vs RadiologyRequestDetail by HISRequestNo + HISTestCode')
    bullet(doc, 'Train users: radiology workflow starts at Radiology Report Entry queue, not Sample Collection')

    h(doc, '17. Conclusion', 1)
    p(doc, (
        'ZoryaLIS implements radiology as a department-classified extension of the shared HIS Test Master, '
        'with billing through Sale Invoice and operational fulfillment through RadiologyRequestDetail / '
        'RadiologyResultDetail. Laboratory and radiology paths split at SaleInvoiceManager.Save after '
        'SaleInvoiceDetail persistence. Understanding department assignment and the IsRadiologyTest rule '
        'is essential for correct routing. This document reflects the codebase as of June 2026 reverse-engineering '
        'of the avs-lis repository.'
    ))

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    doc.save(OUT)
    print(f'Written: {OUT}')

if __name__ == '__main__':
    main()
